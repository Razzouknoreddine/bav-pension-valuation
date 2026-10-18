"""Wandelt die Original-Excel-Dateien von Destatis und Bundesbank in kleine CSV-Dateien um.

Originaldateien (nicht im Repository, 18 MB) nach data/raw/ legen:
  - statistischer-bericht-kohortensterbetafeln-5126101269005.xlsx   (Destatis, 16.09.2026)
  - abzinsungszinssaetze-10jahre-excel-data.xlsx                     (Deutsche Bundesbank)

Aufruf:  python scripts/prepare_data.py
Ergebnis: data/destatis_kohortensterbetafeln.csv und data/bundesbank_hgb_zins_10j.csv
"""
import csv
import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"

# Tabellenblätter mit qx: b01-b06 = (Variante 1, 2, 3) x (männlich, weiblich)
BLAETTER = {f"csv-12621-b0{i}": None for i in range(1, 7)}
VARIANTE = {"V1 (niedrig)": 1, "V2 (mittel)": 2, "V3 (hoch)": 3}
GESCHLECHT = {"maennlich": "m", "weiblich": "w"}


def kohortentafeln():
    datei = next(RAW.glob("*kohortensterbetafeln*.xlsx"))
    wb = openpyxl.load_workbook(datei, read_only=True)
    ziel = ROOT / "data" / "destatis_kohortensterbetafeln.csv"
    n = 0
    with open(ziel, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["jahrgang", "geschlecht", "variante", "alter", "qx"])
        for blatt in BLAETTER:
            for statistik, jahrgang, geschl, variante, alter, qx in wb[blatt].iter_rows(min_row=2, values_only=True):
                alter = 100 if str(alter).startswith("100") else int(alter)   # "100 und mehr"
                w.writerow([jahrgang, GESCHLECHT[geschl], VARIANTE[variante], alter, f"{float(qx):.8f}"])
                n += 1
    print(f"{ziel.name}: {n} Zeilen")


def bundesbank_zinsen():
    datei = next(RAW.glob("abzinsungszinssaetze-10jahre*.xlsx"))
    zeilen = list(openpyxl.load_workbook(datei, read_only=True).active.iter_rows(values_only=True))
    kopf = next(r for r in zeilen if r[1] == 1)              # Zeile mit Laufzeiten 1 ... 50
    daten = [r for r in zeilen if isinstance(r[0], datetime.datetime)]
    ziel = ROOT / "data" / "bundesbank_hgb_zins_10j.csv"
    with open(ziel, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["stichtag"] + [f"L{j}" for j in kopf[1:]])
        for r in daten:
            w.writerow([r[0].date().isoformat()] + list(r[1:]))
    print(f"{ziel.name}: {len(daten)} Monate ({daten[0][0].date()} bis {daten[-1][0].date()})")


if __name__ == "__main__":
    kohortentafeln()
    bundesbank_zinsen()
